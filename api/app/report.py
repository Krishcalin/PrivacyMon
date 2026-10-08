"""DPIA report rendering (SRS 3.7 / 9): a published DPIA as HTML, PDF or Word.

HTML is pure-Python (always available). PDF (WeasyPrint) and DOCX (python-docx) are
imported lazily inside the renderers, so importing this module — and the whole API —
needs neither library; a deployment without them still serves HTML and returns a clear
501 for the others. Rendering is synchronous here; the SRS's async render-to-object-
store is a later refinement.
"""
from __future__ import annotations

import html
import io
import uuid

from fastapi import APIRouter, HTTPException, Query
from fastapi.responses import HTMLResponse, Response

from dpia_core.controls import QUESTIONNAIRE
from platform_db.models.assessment import (
    DpiaAssessment, QuestionnaireResponse, Risk,
)
from platform_db.models.registry import Application

from . import db
from .settings import settings

router = APIRouter(prefix=settings.api_prefix)
_ANSWER_LABEL = {"yes": "Yes", "partial": "Partial", "no": "No", "na": "N/A"}


def _gather(dpia_id: uuid.UUID) -> dict:
    with db.get_sessionmaker()() as s:
        d = s.get(DpiaAssessment, dpia_id)
        if d is None:
            raise HTTPException(status_code=404, detail="DPIA not found")
        app = s.get(Application, d.application_id)
        responses = {
            r.question_key: r for r in s.query(QuestionnaireResponse).filter(
                QuestionnaireResponse.dpia_id == dpia_id).all()
        }
        risks = s.query(Risk).filter(Risk.dpia_id == dpia_id).order_by(
            Risk.created_at).all()
        sections = []
        for sec in QUESTIONNAIRE:
            qs = []
            for q in sec.questions:
                r = responses.get(q.key)
                qs.append({
                    "key": q.key, "text": q.text,
                    "answer": _ANSWER_LABEL.get(r.answer.value, "—") if r else "Not answered",
                    "justification": (r.justification if r else "") or "",
                })
            sections.append({"key": sec.key, "title": sec.title, "questions": qs})
        return {
            "dpia": d, "app": app, "sections": sections,
            "inventory": d.inventory_snapshot or {},
            "risks": [{
                "title": r.title, "likelihood": r.likelihood, "impact": r.impact,
                "treatment": r.treatment.value if r.treatment else "—",
                "status": r.status.value,
            } for r in risks],
        }


def render_html(data: dict) -> str:
    d, app = data["dpia"], data["app"]
    e = html.escape

    inv_rows = "".join(
        f"<tr><td>{e(cat)}</td><td>{e(inv['tier'])}</td>"
        f"<td style='text-align:right'>{inv.get('locations_count', 0)}</td></tr>"
        for cat, inv in sorted(data["inventory"].items()))

    q_html = ""
    for sec in data["sections"]:
        q_html += f"<h3>{e(sec['key'])}. {e(sec['title'])}</h3><table class='q'>"
        for q in sec["questions"]:
            just = f"<div class='just'>{e(q['justification'])}</div>" if q["justification"] else ""
            q_html += (f"<tr><td class='qt'>{e(q['key'])} — {e(q['text'])}{just}</td>"
                       f"<td class='qa'>{e(q['answer'])}</td></tr>")
        q_html += "</table>"

    risk_rows = "".join(
        f"<tr><td>{e(r['title'])}</td><td style='text-align:center'>{r['likelihood']}×{r['impact']}"
        f"={r['likelihood'] * r['impact']}</td><td>{e(r['treatment'])}</td><td>{e(r['status'])}</td></tr>"
        for r in data["risks"]) or "<tr><td colspan='4' class='muted'>No risks recorded.</td></tr>"

    band = d.risk_band.value if d.risk_band else "—"
    inherent = f"{float(d.inherent_score):.1f}" if d.inherent_score is not None else "—"
    residual = f"{float(d.residual_score):.1f}" if d.residual_score is not None else "—"
    approved = d.approved_at and "Approved" or ""
    return f"""<!doctype html><html><head><meta charset="utf-8"><title>DPIA — {e(app.name if app else '')}</title>
<style>
body{{font-family:'Segoe UI',system-ui,sans-serif;color:#15233f;max-width:860px;margin:0 auto;padding:36px 44px;font-size:13px;}}
h1{{font-size:24px;margin:0 0 2px;}} h2{{font-size:16px;border-bottom:2px solid #1570ef;padding-bottom:4px;margin-top:28px;}}
h3{{font-size:13.5px;margin:16px 0 6px;color:#1570ef;}}
.sub{{color:#5b6b86;margin-bottom:18px;}}
table{{width:100%;border-collapse:collapse;margin:6px 0;}}
th,td{{border:1px solid #d7e0ef;padding:6px 9px;text-align:left;vertical-align:top;}}
th{{background:#f6f9ff;font-size:11px;text-transform:uppercase;color:#5b6b86;}}
.q td.qt{{width:80%;}} .q td.qa{{width:20%;font-weight:600;text-align:center;}}
.just{{color:#5b6b86;font-size:12px;margin-top:3px;}} .muted{{color:#5b6b86;}}
.band{{display:inline-block;padding:3px 12px;border-radius:999px;font-weight:700;}}
.kpi{{display:flex;gap:24px;margin:10px 0;}} .kpi div b{{font-size:22px;display:block;}}
.signoff{{margin-top:40px;border-top:1px solid #d7e0ef;padding-top:16px;}}
</style></head><body>
<h1>Data Protection Impact Assessment</h1>
<div class="sub">{e(app.name if app else '')} · DPDP Act 2023 · Status: {e(d.state.value)}</div>

<h2>1. Application</h2>
<table><tr><th>Environment</th><td>{e(app.environment.value if app else '')}</td>
<th>Hosting</th><td>{e(app.hosting.value if app else '')}</td></tr>
<tr><th>Internet-facing</th><td>{'Yes' if app and app.internet_facing else 'No'}</td>
<th>User base</th><td>{e(app.user_base.value if app and app.user_base else '—')}</td></tr></table>

<h2>2. Personal-data inventory</h2>
<table><tr><th>Category</th><th>Tier</th><th style="text-align:right">Locations</th></tr>{inv_rows or '<tr><td colspan=3 class=muted>No inventory.</td></tr>'}</table>

<h2>3. Risk</h2>
<div class="kpi"><div>Inherent<b>{inherent}</b></div><div>Residual<b>{residual}</b></div>
<div>Band<b><span class="band">{e(band)}</span></b></div></div>

<h2>4. Risk register</h2>
<table><tr><th>Risk</th><th>L×I</th><th>Treatment</th><th>Status</th></tr>{risk_rows}</table>

<h2>5. Questionnaire (A–K)</h2>
{q_html}

<div class="signoff">
<p><b>DPO approval:</b> {e(approved) or 'Pending'}{' · ' + str(d.approved_at) if d.approved_at else ''}</p>
<p><b>Published:</b> {str(d.published_at) if d.published_at else 'Not published'}</p>
<p class="muted">Generated by PrivacyMon. This DPIA reflects the scan-discovered inventory and the owner's attested answers.</p>
</div>
</body></html>"""


def render_pdf(data: dict) -> bytes:
    try:
        from weasyprint import HTML
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=501, detail=f"PDF rendering unavailable: {e}")
    return HTML(string=render_html(data)).write_pdf()


def render_docx(data: dict) -> bytes:
    try:
        from docx import Document
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=501, detail=f"DOCX rendering unavailable: {e}")
    d, app = data["dpia"], data["app"]
    doc = Document()
    doc.add_heading("Data Protection Impact Assessment", 0)
    doc.add_paragraph(f"{app.name if app else ''} · DPDP Act 2023 · Status: {d.state.value}")

    doc.add_heading("Application", level=1)
    t = doc.add_table(rows=0, cols=2)
    for k, v in (("Environment", app.environment.value if app else ""),
                 ("Hosting", app.hosting.value if app else ""),
                 ("Internet-facing", "Yes" if app and app.internet_facing else "No"),
                 ("User base", app.user_base.value if app and app.user_base else "—")):
        row = t.add_row().cells
        row[0].text, row[1].text = k, str(v)

    doc.add_heading("Personal-data inventory", level=1)
    it = doc.add_table(rows=1, cols=3)
    it.rows[0].cells[0].text, it.rows[0].cells[1].text, it.rows[0].cells[2].text = \
        "Category", "Tier", "Locations"
    for cat, inv in sorted(data["inventory"].items()):
        c = it.add_row().cells
        c[0].text, c[1].text, c[2].text = cat, inv["tier"], str(inv.get("locations_count", 0))

    doc.add_heading("Risk", level=1)
    inh = f"{float(d.inherent_score):.1f}" if d.inherent_score is not None else "—"
    res = f"{float(d.residual_score):.1f}" if d.residual_score is not None else "—"
    doc.add_paragraph(f"Inherent: {inh}   Residual: {res}   Band: "
                      f"{d.risk_band.value if d.risk_band else '—'}")

    doc.add_heading("Risk register", level=1)
    for r in data["risks"]:
        doc.add_paragraph(
            f"{r['title']} — L×I {r['likelihood']}×{r['impact']} "
            f"({r['likelihood'] * r['impact']}), {r['treatment']}, {r['status']}",
            style="List Bullet")

    doc.add_heading("Questionnaire (A–K)", level=1)
    for sec in data["sections"]:
        doc.add_heading(f"{sec['key']}. {sec['title']}", level=2)
        for q in sec["questions"]:
            doc.add_paragraph(f"{q['key']} — {q['text']}: {q['answer']}", style="List Bullet")
            if q["justification"]:
                doc.add_paragraph(q["justification"])

    doc.add_heading("Sign-off", level=1)
    doc.add_paragraph(f"DPO approval: {'Approved' if d.approved_at else 'Pending'}")
    doc.add_paragraph(f"Published: {d.published_at or 'Not published'}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@router.get("/dpias/{dpia_id}/report")
def get_report(dpia_id: uuid.UUID, format: str = Query("html", pattern="^(html|pdf|docx)$")):
    data = _gather(dpia_id)
    name = (data["app"].name if data["app"] else "dpia").replace(" ", "_")
    if format == "html":
        return HTMLResponse(render_html(data))
    if format == "pdf":
        return Response(render_pdf(data), media_type="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="{name}_DPIA.pdf"'})
    return Response(
        render_docx(data),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{name}_DPIA.docx"'})
